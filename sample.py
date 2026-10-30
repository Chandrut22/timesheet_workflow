#!/usr/bin/env python3
"""Builds the Timesheet Assistant test pack from the user's employee list and 9 test cases."""
import csv
import hashlib
import random
import shutil
from datetime import date, time as dtime, timedelta
from pathlib import Path

import cv2
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from PIL import Image, ImageDraw, ImageFont
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfgen import canvas
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

OUT = Path("outputs/timesheet_test_pack")


W0, W1, W2 = date(2026, 10, 5), date(2026, 10, 12), date(2026, 10, 19)
DAYS = ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"]
KT, CM = "Knowledge Tech", "CodingMart"

COMPANIES = {
    KT: dict(aliases="Knowledge Tech.;KnowledgeTech;KT", hourly=350, ot=525, hol=700),
    CM: dict(aliases="Coding Mart;CodingMart Pvt Ltd", hourly=400, ot=600, hol=800),
}
# The user's employee list. Row 5 (Dhaunsh Raj O) reused SYN-EMP-004 in the original file;
# employee_code must be unique per company, so it is SYN-EMP-005 here.
EMPLOYEES = [
    (KT, "SYN-EMP-001", "Chandru Thambidurai"),
    (KT, "SYN-EMP-002", "Deepak"),
    (KT, "SYN-EMP-003", "Saravanan Arunchalam"),
    (KT, "SYN-EMP-004", "S Dhaunsh"),
    (KT, "SYN-EMP-005", "Dhaunsh Raj O"),
    (CM, "CM-EMP-001", "Naveen"),
    (CM, "CM-EMP-002", "Boobala Murgan"),
    (CM, "CM-EMP-003", "HariHaraSudhan"),
    (CM, "CM-EMP-004", "HariPrasath"),
    (CM, "CM-EMP-005", "Boobalan"),
]
EXTRA_CASE3 = [
    (CM, "CM-EMP-006", "Hari Krishnan"),
    (CM, "CM-EMP-007", "Hari Sundar"),
    (CM, "CM-EMP-008", "Hari Vignesh"),
]
HOLIDAYS = [  # (company or "" = all companies, date, name)
    ("", date(2026, 1, 1), "New Year's Day"),
    ("", date(2026, 1, 15), "Pongal"),
    ("", date(2026, 1, 16), "Thiruvalluvar Day"),
    ("", date(2026, 1, 17), "Uzhavar Thirunal"),
    ("", date(2026, 1, 26), "Republic Day"),
    ("", date(2026, 4, 14), "Tamil New Year"),
    ("", date(2026, 5, 1), "May Day"),
    ("", date(2026, 8, 15), "Independence Day"),
    ("", date(2026, 10, 2), "Gandhi Jayanthi"),
    ("", date(2026, 10, 19), "Ayudha Pooja"),
    ("", date(2026, 10, 20), "Vijayadasami"),
    ("", date(2026, 11, 8), "Deepavali"),
    ("", date(2026, 12, 25), "Christmas"),
    (CM, date(2026, 10, 22), "CodingMart Foundation Day (test holiday)"),
    (KT, date(2026, 10, 23), "Knowledge Tech Annual Day (test holiday)"),
]

# ---------------------------------------------------------------- pay rules (mirror of pay_rules.py)


def holiday_dates(company):
    return {d for c, d, _ in HOLIDAYS if c in ("", company)}


def pay(company, monday, hours):
    rate = COMPANIES[company]
    hol = holiday_dates(company)
    per = []
    for i, h in enumerate(hours):
        d = monday + timedelta(i)
        if h <= 0:
            per.append([0.0, 0.0, 0.0])
        elif d in hol:
            per.append([0.0, 0.0, float(h)])
        else:
            per.append([float(min(h, 8)), float(max(h - 8, 0)), 0.0])
    over = max(sum(p[0] for p in per) - 40, 0)
    for p in reversed(per):
        if over <= 0:
            break
        take = min(over, p[0])
        p[0] -= take
        p[1] += take
        over -= take
    reg, ot, hl = (sum(p[k] for p in per) for k in range(3))
    gross = reg * rate["hourly"] + ot * rate["ot"] + hl * rate["hol"]
    return dict(regular=reg, overtime=ot, holiday=hl, total=reg + ot + hl, gross=gross)


def times(h):
    brk = 30 if h >= 5 else 0
    end = 9 * 60 + int(h * 60) + brk
    return dtime(9, 0), dtime(end // 60, end % 60), brk


def write_csv(path, header, rows):
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(header)
        w.writerows(rows)


# ---------------------------------------------------------------- spreadsheet timesheet


def make_xlsx(path, company_text, name_text, code_text, monday, hours):
    wb = Workbook()
    ws = wb.active
    ws.title = "Timesheet"
    base = Font(name="Arial", size=10)
    bold = Font(name="Arial", size=10, bold=True)
    thin = Side(style="thin", color="999999")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    ws["A1"] = "WEEKLY TIMESHEET"
    ws["A1"].font = Font(name="Arial", size=14, bold=True)
    info = [("Company:", company_text), ("Employee:", name_text)]
    if code_text:
        info.append(("Employee ID:", code_text))
    info.append(("Week commencing:", monday))
    r = 3
    for k, v in info:
        ws.cell(r, 1, k).font = bold
        c = ws.cell(r, 2, v)
        c.font = base
        c.alignment = Alignment(horizontal="left")
        if isinstance(v, date):
            c.number_format = "dd-mmm-yyyy"
        r += 1
    head = r + 1
    for j, t in enumerate(["Date", "Day", "Start", "End", "Break (min)", "Hours", "Notes"], 1):
        c = ws.cell(head, j, t)
        c.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F3864")
        c.alignment = Alignment(horizontal="center")
        c.border = box
    for i in range(7):
        row = head + 1 + i
        d = monday + timedelta(i)
        ws.cell(row, 1, d).number_format = "dd-mmm-yyyy"
        ws.cell(row, 2, DAYS[i])
        h = hours[i]
        if h > 0:
            s, e, b = times(h)
            ws.cell(row, 3, s).number_format = "hh:mm"
            ws.cell(row, 4, e).number_format = "hh:mm"
            ws.cell(row, 5, b)
            ws.cell(row, 6, f"=ROUND((D{row}-C{row})*24-E{row}/60,2)")
        for j in range(1, 8):
            c = ws.cell(row, j)
            c.font = base
            c.border = box
            c.alignment = Alignment(horizontal="center")
        ws.cell(row, 6).number_format = "0.00"
    first, last = head + 1, head + 7
    tot = last + 1
    ws.cell(tot, 1, "Total hours").font = bold
    c = ws.cell(tot, 6, f"=SUM(F{first}:F{last})")
    c.font = bold
    c.number_format = "0.00"
    c.alignment = Alignment(horizontal="center")
    for j in range(1, 8):
        ws.cell(tot, j).border = box
    ws.cell(tot + 2, 1, "Employee signature:").font = bold
    ws.cell(tot + 3, 1, "Manager approval:").font = bold
    for col, w in zip("ABCDEFG", [18, 8, 9, 9, 12, 9, 22]):
        ws.column_dimensions[col].width = w
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


# ---------------------------------------------------------------- csv and typed pdf timesheets


def make_csv_timesheet(path, company, name, code, monday, hours):
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["Weekly timesheet"])
        w.writerow(["Company", company])
        w.writerow(["Employee", name])
        w.writerow(["Employee ID", code])
        w.writerow(["Week commencing", monday.isoformat()])
        w.writerow([])
        w.writerow(["Date", "Day", "Hours"])
        for i in range(7):
            w.writerow([(monday + timedelta(i)).isoformat(), DAYS[i], hours[i] if hours[i] else ""])
        w.writerow(["Total", "", sum(hours)])


def make_pdf(path, company, name, code, monday, hours, stated_total=None):
    path.parent.mkdir(parents=True, exist_ok=True)
    doc = SimpleDocTemplate(str(path), pagesize=A4, leftMargin=20 * mm, rightMargin=20 * mm,
                            topMargin=20 * mm, bottomMargin=20 * mm, title="Weekly timesheet",
                            author="Test data")
    ss = getSampleStyleSheet()
    story = [Paragraph("WEEKLY TIMESHEET", ss["Title"]), Spacer(1, 8)]
    info = [["Company:", company], ["Employee:", name]]
    if code:
        info.append(["Employee ID:", code])
    info.append(["Week commencing:", monday.strftime("%d-%b-%Y")])
    t = Table(info, colWidths=[42 * mm, 110 * mm])
    t.setStyle(TableStyle([("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
                           ("FONTNAME", (1, 0), (1, -1), "Helvetica"), ("FONTSIZE", (0, 0), (-1, -1), 10)]))
    story += [t, Spacer(1, 12)]
    rows = [["Date", "Day", "Start", "End", "Break (min)", "Hours"]]
    for i in range(7):
        d = (monday + timedelta(i)).strftime("%d-%b-%Y")
        h = hours[i]
        if h > 0:
            s, e, b = times(h)
            rows.append([d, DAYS[i], s.strftime("%H:%M"), e.strftime("%H:%M"), str(b), f"{h:.1f}"])
        else:
            rows.append([d, DAYS[i], "", "", "", ""])
    total = sum(hours) if stated_total is None else stated_total
    rows.append(["Total hours", "", "", "", "", f"{total:.1f}"])
    tb = Table(rows, colWidths=[34 * mm, 18 * mm, 24 * mm, 24 * mm, 28 * mm, 24 * mm])
    tb.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey), ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F3864")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"), ("ALIGN", (1, 0), (-1, -1), "CENTER"),
        ("FONTSIZE", (0, 0), (-1, -1), 10), ("TOPPADDING", (0, 0), (-1, -1), 6), ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
    ]))
    story += [tb, Spacer(1, 28), Paragraph("Employee signature: ______________________", ss["Normal"]),
              Spacer(1, 14), Paragraph("Manager approval: ______________________", ss["Normal"])]
    doc.build(story)


# ---------------------------------------------------------------- scanned (image-only) pdf


def make_scanned_pdf(path, company, name, code, monday, hours):
    W, H = 1654, 2339  # A4 at 200 dpi
    img = Image.new("RGB", (W, H), "white")
    d = ImageDraw.Draw(img)
    ft = ImageFont.load_default(size=72)
    fb = ImageFont.load_default(size=40)
    f = ImageFont.load_default(size=40)
    d.text((W // 2, 150), "WEEKLY TIMESHEET", font=ft, fill="black", anchor="mm")
    y = 280
    lines = [("Company:", company), ("Employee:", name), ("Employee ID:", code),
             ("Week commencing:", monday.strftime("%d-%b-%Y"))]
    for k, v in lines:
        d.text((150, y), k, font=fb, fill="black")
        d.text((520, y), v, font=f, fill="black")
        y += 62
    cols = [150, 520, 740, 960, 1180, 1400]
    heads = ["Date", "Day", "Start", "End", "Break", "Hours"]
    top = y + 50
    row_h = 84
    d.rectangle([cols[0], top, 1504, top + row_h], fill=(225, 225, 225))
    for x, t in zip(cols, heads):
        d.text((x + 14, top + 20), t, font=fb, fill="black")
    for i in range(7):
        ry = top + row_h * (i + 1)
        h = hours[i]
        vals = [(monday + timedelta(i)).strftime("%d-%b"), DAYS[i], "", "", "", ""]
        if h > 0:
            s, e, b = times(h)
            vals = [vals[0], DAYS[i], s.strftime("%H:%M"), e.strftime("%H:%M"), str(b), f"{h:.1f}"]
        for x, t in zip(cols, vals):
            d.text((x + 14, ry + 20), t, font=f, fill="black")
    ry = top + row_h * 8
    d.text((cols[0] + 14, ry + 20), "Total hours", font=fb, fill="black")
    d.text((cols[5] + 14, ry + 20), f"{sum(hours):.1f}", font=fb, fill="black")
    for i in range(10):
        yy = top + row_h * i
        d.line([cols[0], yy, 1504, yy], fill="black", width=2)
    for x in cols + [1504]:
        d.line([x, top, x, top + row_h * 9], fill="black", width=2)
    d.text((150, ry + 200), "Employee signature: ____________________", font=f, fill="black")
    d.text((150, ry + 290), "Manager approval:   ____________________", font=f, fill="black")

    rng = np.random.default_rng(3)
    a = np.array(img.convert("L")).astype(np.float32)
    yy, xx = np.mgrid[0:H, 0:W]
    a = a * (1 - 0.07 * (xx / W) - 0.05 * (yy / H))
    a = cv2.GaussianBlur(a, (0, 0), 1.2) + rng.normal(0, 4, a.shape)
    M = cv2.getRotationMatrix2D((W / 2, H / 2), 0.7, 1.0)
    a = cv2.warpAffine(a, M, (W, H), borderValue=205)
    a[:, :5] = 80
    a = np.clip(a, 0, 255).astype(np.uint8)
    tmp = path.with_suffix(".tmp.jpg")
    path.parent.mkdir(parents=True, exist_ok=True)
    cv2.imwrite(str(tmp), a, [cv2.IMWRITE_JPEG_QUALITY, 72])
    c = canvas.Canvas(str(path), pagesize=A4)
    c.drawImage(str(tmp), 0, 0, width=A4[0], height=A4[1])
    c.save()
    tmp.unlink()


# ---------------------------------------------------------------- handwritten photo


def make_handwritten_photo(path, hours):
    rng = random.Random(11)
    nrng = np.random.default_rng(11)
    W, H = 1500, 2000
    paper = np.full((H, W, 3), (236, 239, 241), np.float32)
    for y in range(160, H - 60, 58):
        cv2.line(paper, (0, y), (W, y), (232, 214, 190), 2)
    cv2.line(paper, (110, 0), (110, H), (190, 190, 232), 3)
    ink = np.full((H, W, 3), 255, np.uint8)
    BLUE, PEN = (130, 55, 25), (90, 60, 45)
    SC = cv2.FONT_HERSHEY_SCRIPT_COMPLEX
    SS = cv2.FONT_HERSHEY_SCRIPT_SIMPLEX

    def text(s, x, y, scale, thick=3, font=SC, color=BLUE):
        for ch in s:
            if ch == " ":
                x += int(30 * scale)
                continue
            sc = scale * rng.uniform(0.94, 1.08)
            t = thick + (1 if rng.random() < 0.2 else 0)
            cv2.putText(ink, ch, (int(x), int(y + rng.randint(-3, 3))), font, sc, color, t, cv2.LINE_AA)
            x += cv2.getTextSize(ch, font, sc, t)[0][0] + rng.randint(2, 6)
        return x

    def line(p1, p2, thick=3):
        n = max(int(np.hypot(p2[0] - p1[0], p2[1] - p1[1]) // 60), 2)
        pts = [(p1[0] + (p2[0] - p1[0]) * i / n + rng.uniform(-2.5, 2.5),
                p1[1] + (p2[1] - p1[1]) * i / n + rng.uniform(-2.5, 2.5)) for i in range(n + 1)]
        cv2.polylines(ink, [np.array(pts, np.int32)], False, PEN, thick, cv2.LINE_AA)

    text("Coding Mart", 150, 200, 2.6, 4)
    text("Weekly Timesheet", 150, 300, 1.6, 3)
    text("Name: Naveen", 150, 420, 1.6, 3)
    text("ID: CM-EMP-001", 800, 420, 1.6, 3)
    text("Week: 12 Oct - 18 Oct 2026", 150, 520, 1.6, 3)
    xs = [110, 380, 620, 860, 1100, 1400]
    ys = [590, 700, 820, 940, 1060, 1180, 1300, 1430]
    heads = ["Date", "Day", "In", "Out", "Hours"]
    for x, t in zip(xs, heads):
        text(t, x + 40, 665, 1.4, 3)
    ins, outs = "09:00", None
    for i in range(5):
        h = hours[i]
        s, e, b = times(h)
        row = [f"{12 + i}/10", DAYS[i], "09:00", e.strftime("%H:%M"), f"{h:g}"]
        for x, t in zip(xs, row):
            text(t, x + 40, ys[i + 1] - 38, 1.5, 3, SS if t[:1].isdigit() else SC)
    text("Total", xs[0] + 40, ys[6] - 38, 1.5, 3)
    text(f"{sum(hours):g}", xs[4] + 40, ys[6] - 38, 1.6, 4, SS)
    for y in ys[:7]:
        line((xs[0], y), (xs[5], y))
    line((xs[0], ys[6] + 8), (xs[5], ys[6] + 8))
    for x in xs:
        line((x, ys[0]), (x, ys[6]))
    text("Signed:", 150, 1560, 1.5, 3)
    text("Naveen", 420, 1565, 2.0, 3)

    paper = paper * (ink.astype(np.float32) / 255.0)
    paper += nrng.normal(0, 2.5, paper.shape)
    ow, oh = W + 260, H + 260
    desk = np.zeros((oh, ow, 3), np.float32)
    desk[:] = (58, 72, 92)
    desk += nrng.normal(0, 6, desk.shape)
    src = np.float32([[0, 0], [W, 0], [W, H], [0, H]])
    dst = np.float32([[150, 120], [W + 95, 155], [W + 130, H + 150], [110, H + 108]])
    M = cv2.getPerspectiveTransform(src, dst)
    warped = cv2.warpPerspective(np.clip(paper, 0, 255), M, (ow, oh))
    mask = cv2.GaussianBlur(cv2.warpPerspective(np.ones((H, W), np.float32), M, (ow, oh)), (0, 0), 1.5)[..., None]
    yy, xx = np.mgrid[0:oh, 0:ow]
    light = 1.05 - 0.22 * (xx / ow) - 0.10 * (yy / oh) + 0.025 * np.sin(yy / 170.0)
    out = (desk * (1 - mask) + warped * mask) * light[..., None]
    out += nrng.normal(0, 3.0, out.shape)
    out = cv2.GaussianBlur(out, (0, 0), 0.8)
    path.parent.mkdir(parents=True, exist_ok=True)
    cv2.imwrite(str(path), np.clip(out, 0, 255).astype(np.uint8), [cv2.IMWRITE_JPEG_QUALITY, 80])


def make_corrupt_pdf(src_pdf, dst):
    data = Path(src_pdf).read_bytes()
    junk = bytes(random.Random(5).randrange(256) for _ in range(1800))
    Path(dst).write_bytes(data[: int(len(data) * 0.35)] + junk)


# ---------------------------------------------------------------- email text


def email_txt(path, sender, subject, attachments, body):
    path.parent.mkdir(parents=True, exist_ok=True)
    att = "; ".join(attachments) if attachments else "(none: plain text only)"
    path.write_text(
        f"FROM:     {sender}\nTO:       <the Gmail inbox the app monitors>\nSUBJECT:  {subject}\n"
        f"ATTACH:   {att}\n{'-' * 64}\n{body.strip()}\n", encoding="utf-8")


# ================================================================ build
if OUT.exists():
    shutil.rmtree(OUT)
OUT.mkdir(parents=True)
EXP = []


def expect(case, file, company, emp, code, monday, hours, status, flags, note):
    p = pay(company, monday, hours)
    EXP.append([case, file, company, emp, code, monday.isoformat(), f"{p['regular']:g}", f"{p['overtime']:g}",
                f"{p['holiday']:g}", f"{p['total']:g}", f"{p['gross']:.2f}", status, flags, note])


# ---- setup files
S = OUT / "setup"
write_csv(S / "companies.csv",
          ["name", "aliases", "hourly_rate", "overtime_rate", "double_time_rate", "holiday_rate", "currency",
           "daily_ot_after", "weekly_ot_after", "double_time_after", "week_start", "is_active"],
          [[n, c["aliases"], c["hourly"], c["ot"], "", c["hol"], "INR", 8, 40, "", 0, "true"]
           for n, c in COMPANIES.items()])
EMP_HEAD = ["company", "employee_code", "full_name", "email", "aliases", "assignment_start", "assignment_end"]
write_csv(S / "employees.csv", EMP_HEAD, [[c, k, n, "", "", "", ""] for c, k, n in EMPLOYEES])
write_csv(S / "employees_extra_case3.csv", EMP_HEAD, [[c, k, n, "", "", "", ""] for c, k, n in EXTRA_CASE3])
write_csv(S / "employees_update_case9.csv", EMP_HEAD, [[CM, "CM-EMP-006", "Hari Krishnan", "", "Hari", "", ""]])
write_csv(S / "holidays.csv", ["company", "holiday_date", "name"],
          [[c, d.isoformat(), n] for c, d, n in HOLIDAYS])
wb = Workbook()
ws = wb.active
ws.title = "Holidays"
for j, t in enumerate(["company", "holiday_date", "name"], 1):
    c = ws.cell(1, j, t)
    c.font = Font(name="Arial", size=10, bold=True)
for i, (c, d, n) in enumerate(HOLIDAYS, 2):
    ws.cell(i, 1, c or None)
    ws.cell(i, 2, d).number_format = "yyyy-mm-dd"
    ws.cell(i, 3, n)
    for j in range(1, 4):
        ws.cell(i, j).font = Font(name="Arial", size=10)
for col, w in zip("ABC", [18, 14, 44]):
    ws.column_dimensions[col].width = w
note = wb.create_sheet("Notes")
for i, t in enumerate([
    "Import reads the first sheet only (Holidays). This sheet is for people.",
    "A blank company means the holiday applies to every company.",
    "Dates 2026-10-19 and 2026-10-20: Tamil Nadu notification as reported (Ayudha Pooja, Vijayadasami).",
    "The CodingMart (22 Oct) and Knowledge Tech (23 Oct) rows are invented test holidays for case 7.",
    "Check the national and state dates against your HR calendar before using them for real payroll."], 1):
    note.cell(i, 1, t).font = Font(name="Arial", size=10)
note.column_dimensions["A"].width = 110
wb.save(S / "holidays.xlsx")

# ---- case 1: Excel, initial + surname
c1 = OUT / "case1_excel_initial_surname"
f = "Timesheet_C_Thambidurai_2026-10-12.xlsx"
H1 = [8, 8, 8, 8, 8, 0, 0]
make_xlsx(c1 / f, "Knowledge Tech.", "C. Thambidurai", "", W1, H1)
email_txt(c1 / "email.txt", "<address whitelisted for Knowledge Tech>",
          "Timesheet - C. Thambidurai - week of 12 Oct 2026", [f],
          "Hi,\n\nPlease find the timesheet for C. Thambidurai (Knowledge Tech.) for the week starting 12 Oct 2026.\n\nThanks,\nKnowledge Tech Payroll")
expect(1, f, KT, "Chandru Thambidurai", "SYN-EMP-001", W1, H1, "auto-approved",
       "none", "Matched by initial + surname 'C. Thambidurai' (confidence 0.85)")

# ---- case 2: handwritten photo
c2 = OUT / "case2_handwritten_photo"
f = "Timesheet_Naveen_handwritten_2026-10-12.jpg"
H2 = [8, 10, 6, 8, 8, 0, 0]
make_handwritten_photo(c2 / f, H2)
email_txt(c2 / "email.txt", "<address whitelisted for CodingMart>",
          "Naveen - handwritten timesheet (photo)", [f],
          "Hi,\n\nPhoto of Naveen's handwritten timesheet for 12-18 Oct 2026 attached.\n\nRegards,\nCodingMart Admin")
expect(2, f, CM, "Naveen", "CM-EMP-001", W1, H2, "matched; needs review (handwriting)",
       "none required", "Tuesday 10 h = 8 regular + 2 overtime; the week is 38 regular + 2 overtime")

# ---- case 3: email text only, three Haris
c3 = OUT / "case3_email_text_only_three_haris"
email_txt(c3 / "email.txt", "<address whitelisted for CodingMart>", "Hours for last week", [],
          "Hi,\n\nHere are my hours for the week 12 Oct to 18 Oct 2026:\n\nMon 12 Oct - 8 hours\nTue 13 Oct - 8 hours\n"
          "Wed 14 Oct - 8 hours\nThu 15 Oct - 8 hours\nFri 16 Oct - 8 hours\n\nTotal: 40 hours.\n\nThanks,\nHari")
expect(3, "(email text only)", CM, "(unmatched)", "", W1, [8, 8, 8, 8, 8, 0, 0], "needs review",
       "employee_unmatched", "3 employees fit 'Hari': Hari Krishnan, Hari Sundar, Hari Vignesh (import employees_extra_case3.csv first)")

# ---- case 4: email says 40, PDF says 38
c4 = OUT / "case4_email_vs_pdf_total_mismatch"
f = "Timesheet_Deepak_2026-10-12.pdf"
H4 = [8, 8, 6, 8, 8, 0, 0]
make_pdf(c4 / f, KT, "Deepak", "SYN-EMP-002", W1, H4)
email_txt(c4 / "email.txt", "<address whitelisted for Knowledge Tech>",
          "Deepak - timesheet week of 12 Oct", [f],
          "Hi,\n\nPlease find Deepak's timesheet for the week of 12 Oct 2026 attached.\nTotal hours for the week: 40.\n\nRegards,\nKnowledge Tech Payroll")
expect(4, f, KT, "Deepak", "SYN-EMP-002", W1, H4, "needs review", "total_mismatch",
       "Email says 40 h, the PDF says 38 h; both values shown")

# ---- case 5: three timesheets, one email
c5 = OUT / "case5_three_timesheets_one_email"
f5 = ["01_HariHaraSudhan_2026-10-12.xlsx", "02_HariPrasath_2026-10-12.pdf", "03_Boobalan_2026-10-12.csv"]
H5 = [[8, 8, 8, 8, 8, 0, 0], [8, 8, 8, 8, 4, 0, 0], [9, 8, 8, 8, 7, 0, 0]]
make_xlsx(c5 / f5[0], CM, "HariHaraSudhan", "CM-EMP-003", W1, H5[0])
make_pdf(c5 / f5[1], CM, "HariPrasath", "CM-EMP-004", W1, H5[1])
make_csv_timesheet(c5 / f5[2], CM, "Boobalan", "CM-EMP-005", W1, H5[2])
email_txt(c5 / "email.txt", "<address whitelisted for CodingMart>",
          "Timesheets - week of 12 Oct 2026 (3 people)", f5,
          "Hi,\n\nAttached are the timesheets for HariHaraSudhan, HariPrasath and Boobalan for the week of 12 Oct 2026.\n\nRegards,\nCodingMart Admin")
for fl, nm, cd, hh in zip(f5, ["HariHaraSudhan", "HariPrasath", "Boobalan"],
                          ["CM-EMP-003", "CM-EMP-004", "CM-EMP-005"], H5):
    expect(5, fl, CM, nm, cd, W1, hh, "matched (auto-approve possible)", "none", "One of three separate entries from the same email")

# ---- case 6: same scanned timesheet twice
c6 = OUT / "case6_duplicate_scanned_timesheet"
f = "Scan_Saravanan_2026-10-12.pdf"
f_re = "Scan_Saravanan_2026-10-12_resend.pdf"
H6 = [8, 8, 8, 8, 8, 0, 0]
make_scanned_pdf(c6 / f, KT, "Saravanan Arunchalam", "SYN-EMP-003", W1, H6)
shutil.copyfile(c6 / f, c6 / f_re)
email_txt(c6 / "email_1_first_send.txt", "<address whitelisted for Knowledge Tech>",
          "Scanned timesheet - Saravanan Arunchalam", [f],
          "Hi,\n\nScanned timesheet for Saravanan Arunchalam, week of 12 Oct 2026, attached.\n\nRegards,\nKnowledge Tech Payroll")
email_txt(c6 / "email_2_resend.txt", "<address whitelisted for Knowledge Tech>",
          "Scanned timesheet - Saravanan Arunchalam (sending again)", [f_re],
          "Hi,\n\nSending the same scan again in case the first email did not reach you.\n\nRegards,\nKnowledge Tech Payroll")
expect(6, f, KT, "Saravanan Arunchalam", "SYN-EMP-003", W1, H6, "first send: normal", "none", "First send of the scan")
expect(6, f_re, KT, "Saravanan Arunchalam", "SYN-EMP-003", W1, H6, "flagged duplicate",
       "duplicate_file (and duplicate_week)", "Byte-identical to the first file; do not re-save it")

# ---- case 7: hours on public holidays
c7 = OUT / "case7_hours_on_public_holiday"
f_cm, f_kt = "Timesheet_BoobalaMurgan_2026-10-19.xlsx", "Timesheet_Saravanan_2026-10-19.pdf"
H7 = [8, 8, 8, 8, 8, 0, 0]
make_xlsx(c7 / f_cm, CM, "Boobala Murgan", "CM-EMP-002", W2, H7)
make_pdf(c7 / f_kt, KT, "Saravanan Arunchalam", "SYN-EMP-003", W2, H7)
email_txt(c7 / "email_codingmart.txt", "<address whitelisted for CodingMart>",
          "Timesheet - Boobala Murgan - week of 19 Oct 2026", [f_cm],
          "Hi,\n\nTimesheet for Boobala Murgan for the week of 19 Oct 2026 attached.\n\nRegards,\nCodingMart Admin")
email_txt(c7 / "email_knowledgetech.txt", "<address whitelisted for Knowledge Tech>",
          "Timesheet - Saravanan Arunchalam - week of 19 Oct 2026", [f_kt],
          "Hi,\n\nTimesheet for Saravanan Arunchalam for the week of 19 Oct 2026 attached.\n\nRegards,\nKnowledge Tech Payroll")
expect(7, f_cm, CM, "Boobala Murgan", "CM-EMP-002", W2, H7, "matched", "holiday_worked (non-blocking)",
       "Holiday pay on Mon 19, Tue 20 and Thu 22 (CodingMart-only holiday); Fri 23 is a normal day")
expect(7, f_kt, KT, "Saravanan Arunchalam", "SYN-EMP-003", W2, H7, "matched", "holiday_worked (non-blocking)",
       "Holiday pay on Mon 19, Tue 20 and Fri 23 (Knowledge Tech-only holiday); Thu 22 is a normal day")

# ---- case 8: corrupted attachment in the middle of a batch
c8 = OUT / "case8_corrupted_attachment_in_batch"
f8 = ["01_Deepak_2026-10-05.pdf", "02_Chandru_Thambidurai_2026-10-05.pdf", "03_S_Dhaunsh_2026-10-05.xlsx"]
H8 = [[8, 8, 8, 8, 8, 0, 0], [8, 8, 8, 8, 8, 0, 0], [8, 8, 8, 8, 6, 0, 0]]
make_pdf(c8 / f8[0], KT, "Deepak", "SYN-EMP-002", W0, H8[0])
make_pdf(Path("/tmp/_good_chandru.pdf"), KT, "Chandru Thambidurai", "SYN-EMP-001", W0, H8[1])
make_corrupt_pdf("/tmp/_good_chandru.pdf", c8 / f8[1])
make_xlsx(c8 / f8[2], KT, "S Dhaunsh", "SYN-EMP-004", W0, H8[2])
email_txt(c8 / "email.txt", "<address whitelisted for Knowledge Tech>",
          "Weekly timesheets - 5 to 9 Oct 2026 (3 files)", f8,
          "Hi,\n\nTimesheets for the week of 5 Oct 2026 attached (three files).\n\nRegards,\nKnowledge Tech Payroll")
expect(8, f8[0], KT, "Deepak", "SYN-EMP-002", W0, H8[0], "processed", "none", "Good file before the corrupt one")
EXP.append([8, f8[1], KT, "Chandru Thambidurai", "SYN-EMP-001", W0.isoformat(), "", "", "", "", "",
            "read error; email marked Failed", "(no timesheet)",
            "Corrupted on purpose: clear error shown and logged against the Read step"])
expect(8, f8[2], KT, "S Dhaunsh", "SYN-EMP-004", W0, H8[2], "processed", "none", "Good file after the corrupt one: the batch continues")

# ---- case 9: reprocess after the employee list is updated
c9 = OUT / "case9_reprocess_after_employee_update"
c9.mkdir(parents=True, exist_ok=True)
shutil.copyfile(S / "employees_update_case9.csv", c9 / "employees_update_case9.csv")
expect(9, "(case 3 submission, reprocessed)", CM, "Hari Krishnan", "CM-EMP-006", W1, [8, 8, 8, 8, 8, 0, 0],
       "matched after reprocess", "employee_unmatched cleared",
       "Alias 'Hari' now belongs to Hari Krishnan. Reason entered, reprocess run and history saved")

write_csv(OUT / "expected_results.csv",
          ["case", "file", "company", "employee", "employee_code", "week_start", "regular_hours", "overtime_hours",
           "holiday_hours", "total_hours", "gross_pay_inr", "expected_status", "expected_flags", "notes"], EXP)
print("built", OUT)